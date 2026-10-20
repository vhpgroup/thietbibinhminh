import sqlite3, json
from openpyxl import load_workbook
c=sqlite3.connect('file:C:/Users/mrduy/Downloads/dat_cang_catalog.sqlite?mode=ro',uri=True)
tables=c.execute("SELECT name,sql FROM sqlite_master WHERE type='table'").fetchall()
for name,sql in tables:
 print(name,sql)
 print('COUNT',c.execute('SELECT count(*) FROM "'+name+'"').fetchone()[0])
 print('SAMPLE',json.dumps(c.execute('SELECT * FROM "'+name+'" LIMIT 1').fetchall(),ensure_ascii=False,default=str)[:2500])
w=load_workbook('C:/Users/mrduy/Downloads/products_complete.xlsx',read_only=True,data_only=True)
for s in w:
 print('SHEET',s.title,s.max_row,s.max_column)
 print('HEADERS',json.dumps(list(s.values)[:2],ensure_ascii=False,default=str)[:3500])
print('KINDS',c.execute('SELECT kind,count(*) FROM categories GROUP BY kind').fetchall())
print('ROOTS',c.execute('SELECT id,name,kind FROM categories WHERE parent_id IS NULL').fetchall())
print('PARENT SAMPLE',c.execute('SELECT id,parent_id,name,kind,url FROM categories LIMIT 35').fetchall())
print('ATTR COUNTS',c.execute('SELECT a.name,count(*) FROM product_attribute_values v JOIN attributes a ON a.id=v.attribute_id GROUP BY a.id ORDER BY count(*) DESC LIMIT 12').fetchall())
sheet=w.worksheets[0]
excel=list(sheet.values)
print('EXCEL PRODUCTS',len(excel)-1,'unique SKU',len(set(row[0] for row in excel[1:])))
print('CATEGORY ASSIGNMENT',c.execute('SELECT p.id,c.id,c.name,c.kind FROM products p JOIN product_categories pc ON pc.product_id=p.id JOIN categories c ON c.id=pc.category_id WHERE p.id<4').fetchall())
